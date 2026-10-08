"""Unit and integration tests for TKBI audit checklist generation."""

import pytest
from pathlib import Path
import openpyxl

from sustainmetric.server import generate_tkbi_audit_checklist
from sustainmetric.data.tkbi_sectors_catalog import TKBI_8_SECTORS, map_emiten_to_sectors


def test_tkbi_8_sectors_catalog_structure():
    assert len(TKBI_8_SECTORS) == 8
    required_sectors = [
        "Energi",
        "Konstruksi dan Real Estat",
        "Transportasi dan Pergudangan",
        "Pertanian, Kehutanan, dan Perikanan",
        "Manufaktur",
        "Pengelolaan Air, Air Limbah, Sampah, dan Remediasi",
        "Informasi dan Komunikasi",
        "Aktivitas Profesional, Ilmiah, dan Teknis",
    ]
    for s in required_sectors:
        assert s in TKBI_8_SECTORS
        assert len(TKBI_8_SECTORS[s]["items"]) > 0


def test_sector_mapping():
    # Test geothermal maps to Energi
    mapped = map_emiten_to_sectors("Renewable Energy", "Geothermal power plant developer")
    assert "Energi" in mapped

    # Test coal mining maps to Energi
    mapped_coal = map_emiten_to_sectors("Coal", "Thermal coal mining operations")
    assert "Energi" in mapped_coal

    # Test property maps to Konstruksi dan Real Estat
    mapped_prop = map_emiten_to_sectors("Property", "Gedung apartemen dan real estate")
    assert "Konstruksi dan Real Estat" in mapped_prop


@pytest.mark.asyncio
async def test_generate_tkbi_audit_checklist_pgeo(tmp_path: Path):
    res = await generate_tkbi_audit_checklist(ticker="PGEO", output_dir=str(tmp_path))
    assert res["status"] == "SUCCESS"
    assert res["ticker"] == "PGEO"
    assert "PGEO_audit_TKBI.xlsx" in res["file_name"]

    file_path = Path(res["file_generated"])
    assert file_path.exists()

    # Load workbook and check sheet rows and headers
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    headers = [ws.cell(1, c).value for c in range(1, 13)]
    assert headers[0] == "Kode Emiten"
    assert headers[4] == "TSCID"
    assert headers[7] == "Jawaban AI"
    assert headers[9] == "Reasoning AI"

    # Verify sheet names including Entity Aggregation Summary
    assert "Sheet1" in wb.sheetnames
    assert "Ringkasan Entitas TKBI" in wb.sheetnames
    ws_sum = wb["Ringkasan Entitas TKBI"]
    assert ws_sum.cell(2, 1).value == "Parameter"
    assert ws_sum.cell(3, 1).value == "Kode Emiten"
    assert ws_sum.cell(3, 2).value == "PGEO"


def test_evaluate_criterion_preliminary_and_unverified_labels():
    from sustainmetric.audit_exporter import evaluate_criterion_for_emiten
    mock_item = {
        "bab": "Pembangkitan Tenaga Listrik",
        "kbli": "35101",
        "tsc_id": "TSC-ENE-01",
        "tsc": "Life-cycle GHG < 100g CO2e/kWh",
        "bentuk_jawaban": ["HIJAU", "TRANSISI", "TIDAK"],
    }
    report = {
        "symbol": "PGEO",
        "overview": {"subsector": "Utilities", "description": "Geothermal power plant"},
        "financials": {},
    }
    eval_res = evaluate_criterion_for_emiten("PGEO", mock_item, report, [], [], consistency_score=75.0)

    assert eval_res["audit_tier"] == "PRELIMINARY_HEURISTIC"
    assert eval_res["verification_status"] == "UNVERIFIED"
    assert "https://sectors.app/company/PGEO" in eval_res["bukti"]
    assert "UNVERIFIED" in eval_res["bukti"]
    assert "[Skrining Awal Heuristik]" in eval_res["reasoning"]
    assert "PRELIMINARY_HEURISTIC" in eval_res["auditor_feedback"]


def test_template_audit_tkbi_contains_all_8_sectors():
    template_path = Path("Template_Audit_TKBI.xlsx")
    assert template_path.exists()
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active
    assert ws.title == "Sheet1"

    found_sectors = set()
    for row in range(2, ws.max_row + 1):
        sec = ws.cell(row, 2).value
        if sec:
            found_sectors.add(sec)

    expected_sectors = {
        "Energi",
        "Konstruksi dan Real Estat",
        "Transportasi dan Pergudangan",
        "Pertanian, Kehutanan, dan Perikanan",
        "Manufaktur",
        "Pengelolaan Air, Air Limbah, Sampah, dan Remediasi",
        "Informasi dan Komunikasi",
        "Aktivitas Profesional, Ilmiah, dan Teknis",
    }
    assert found_sectors == expected_sectors
    assert ws.max_row >= 50


def test_multi_sector_mappings_all_8_sectors():
    test_cases = [
        ("Utilities", "Geothermal power generation", "Energi"),
        ("Property", "Pengembangan apartemen dan perumahan", "Konstruksi dan Real Estat"),
        ("Logistics", "Jasa angkutan truk kontainer dan pergudangan", "Transportasi dan Pergudangan"),
        ("Plantation", "Perkebunan kelapa sawit dan CPO", "Pertanian, Kehutanan, dan Perikanan"),
        ("Smelter", "Pabrik pemurnian nikel dan baja", "Manufaktur"),
        ("Sanitasi", "Pengolahan air limbah dan daur ulang sampah", "Pengelolaan Air, Air Limbah, Sampah, dan Remediasi"),
        ("Telecommunication", "Pusat data cloud dan infrastruktur telekomunikasi", "Informasi dan Komunikasi"),
        ("Consulting", "Jasa audit energi dan riset rekayasa", "Aktivitas Profesional, Ilmiah, dan Teknis"),
    ]
    for sub, desc, expected_sec in test_cases:
        res = map_emiten_to_sectors(sub, desc)
        assert expected_sec in res


@pytest.mark.asyncio
async def test_generate_tkbi_audit_checklist_sector_override(tmp_path: Path):
    # Test generating with explicit sector override for Manufaktur
    res_man = await generate_tkbi_audit_checklist(ticker="PGEO", sector="Manufaktur", output_dir=str(tmp_path))
    assert res_man["status"] == "SUCCESS"
    wb_man = openpyxl.load_workbook(res_man["file_generated"])
    ws_man = wb_man.active
    sectors_in_file = {ws_man.cell(r, 2).value for r in range(2, ws_man.max_row + 1)}
    assert "Manufaktur" in sectors_in_file

    # Test generating with sector='all'
    res_all = await generate_tkbi_audit_checklist(ticker="PGEO", sector="all", output_dir=str(tmp_path))
    assert res_all["status"] == "SUCCESS"
    wb_all = openpyxl.load_workbook(res_all["file_generated"])
    ws_all = wb_all.active
    all_sectors_in_file = {ws_all.cell(r, 2).value for r in range(2, ws_all.max_row + 1)}
    assert len(all_sectors_in_file) == 8

